# package

#numpy
#pendulum
#python imaging library
#moviepy
#requests
#thinter
#pyqt
#pandas
#py32
#pytest

"""import camelcase
x = camelcase.CamelCase()
a = "hi i am madhan"
print(x.hump(a))"""
# xlrd is xls
# openpyxl is xlsx
#import xlrd
import openpyxl
print(openpyxl.__version__)
loc="D:\\download\\Sales_Data_50.xlsx"
wb = openpyxl.load_workbook(loc)
sheet = wb.active
print(sheet.max_row)
print(sheet.max_column)
for i in range(1, sheet.max_row + 1):
    print(sheet.cell(i, 1).value)
print([cell.value for cell in sheet["B"]])
"""
import xlrd

loc = "D:\\download\\Sales_Data_50.xlsx"

wb = xlrd.open_workbook(loc)
sheet = wb.sheet_by_index(0)

print(sheet.nrows)
print(sheet.ncols)"""
